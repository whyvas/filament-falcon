import json, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from items import I
V=json.load(open('verified.json'))
fil=json.load(open('filament_estimate.json'))
DATE='2026-10-08'
CUR='"$"#,##0.00'
wb=Workbook(); ws=wb.active; ws.title='BOM'
H=['Category','Item','Spec/notes','Used for','Qty needed','Have it','Qty to buy','Unit price CAD','Line total','Store','Link','Search link','Price checked date','Price status','Cart #']
ws.append(H)
hdr=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='1F3864')
for c in ws[1]: c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
def amz_s(t): return 'https://www.amazon.ca/s?k='+urllib.parse.quote_plus(t)
def kent_s(t): return 'https://kent.ca/en/search?q='+urllib.parse.quote_plus(t)
def hd_s(t): return 'https://www.homedepot.ca/search?q='+urllib.parse.quote_plus(t)
stats=dict(verified=0,search=0,kent=0,est=0,owned=0)
n=0
for (cat,item,spec,used,qty,have,store,ref,est,terms) in I:
    r=ws.max_row+1; n+=1
    link=''; search=''; price=None; date=''; status=''; st=store
    if store=='Owned':
        price=0; status='Owned'; stats['owned']+=1
    elif store=='Amazon.ca':
        v=V.get(ref,{})
        link=f'https://www.amazon.ca/dp/{ref}'
        search=amz_s(terms)
        if v.get('title'):
            spec=spec+f' | Listing: "{v["title"][:100]}"'
            stats['verified']+=1
            if v.get('price'):
                price=v['price']; date=DATE; status='Verified amazon.ca price'
            else:
                price=est; date=DATE; status='Link verified; price NOT shown (placeholder estimate)'
        else:
            link=''; price=est; status='Not verified - use search link (estimate)'; stats['search']+=1
    elif store=='Amazon.ca*':
        st='Amazon.ca'; search=amz_s(terms); price=est; status='No suitable listing found - estimate'; stats['search']+=1
    elif store=='Kent':
        link=ref; search=kent_s(terms); price=est; date=DATE; status='Kent.ca web price (Edmundston store shown)'; stats['kent']+=1
    elif store=='Kent*':
        st='Kent / Home Depot'; search=kent_s(terms); price=est; status='Estimate - check in store'; stats['est']+=1
    elif store=='Local*':
        st='Local metal supplier'; search='https://www.google.com/search?q='+urllib.parse.quote_plus(terms+' Moncton NB'); price=est; status='Estimate - get a quote'; stats['est']+=1
    elif store=='Industrial*':
        st='Industrial supplier'; search=amz_s(terms); price=est; status='Estimate - get a quote'; stats['est']+=1
    if store=='Kent': st='Kent'
    ws.append([cat,item,spec,used,qty,have,None,price,None,st,link,search,date,status,None])
    ws.cell(r,7).value=f'=IF(OR(F{r}=TRUE(),F{r}="TRUE"),0,E{r})'
    ws.cell(r,9).value=f'=IF(H{r}="","",G{r}*H{r})'
    ws.cell(r,15).value=f'=IF(AND(J{r}="Amazon.ca",G{r}>0,ISERROR(SEARCH("ALTERNATIVE",A{r}))),MAX(O$1:O{r-1})+1,"")'
    for col in (11,12):
        c=ws.cell(r,col)
        if c.value: c.hyperlink=c.value; c.style='Hyperlink'
last=ws.max_row
for row in ws.iter_rows(min_row=2,max_row=last):
    for c in row: c.alignment=Alignment(wrap_text=(c.column in (2,3,4,14)),vertical='top')
    row[7].number_format=CUR; row[8].number_format=CUR
dv=DataValidation(type='list',formula1='"TRUE,FALSE"',allow_blank=False,showDropDown=False)
dv.error='Choose TRUE or FALSE'; dv.prompt='TRUE = you already have it'
ws.add_data_validation(dv); dv.add(f'F2:F{last}')
ws.conditional_formatting.add(f'A2:N{last}',FormulaRule(formula=[f'OR($F2=TRUE(),$F2="TRUE")'],fill=PatternFill('solid',fgColor='E2EFDA')))
ws.conditional_formatting.add(f'A2:N{last}',FormulaRule(formula=['ISNUMBER(SEARCH("ALTERNATIVE",$A2))'],fill=PatternFill('solid',fgColor='FFF2CC')))
W=[30,46,70,36,9,9,9,12,12,16,38,38,13,30,7]
for i,w in enumerate(W): ws.column_dimensions[chr(65+i)].width=w
ws.freeze_panes='C2'; ws.auto_filter.ref=f'A1:O{last}'
ws.column_dimensions['O'].hidden=True
# Summary
cats=[]
for r in I:
    if r[0] not in cats: cats.append(r[0])
s=wb.create_sheet('Summary')
s.append(['Category','Lines','Lines still to buy','Total to buy CAD']); 
for c in s[1]: c.font=hdr; c.fill=fill
for i,cat in enumerate(cats):
    r=i+2
    s.cell(r,1,cat)
    s.cell(r,2,f'=COUNTIF(BOM!$A$2:$A${last},A{r})')
    s.cell(r,3,f'=COUNTIFS(BOM!$A$2:$A${last},A{r},BOM!$G$2:$G${last},">0")')
    s.cell(r,4,f'=SUMIF(BOM!$A$2:$A${last},A{r},BOM!$I$2:$I${last})').number_format=CUR
e=len(cats)+1; g=e+2
s.cell(g,1,'GRAND TOTAL to buy (excl. ALTERNATIVE rows)').font=Font(bold=True)
s.cell(g,3,f'=SUMPRODUCT((ISERROR(SEARCH("ALTERNATIVE",A2:A{e})))*C2:C{e})')
s.cell(g,4,f'=SUMPRODUCT((ISERROR(SEARCH("ALTERNATIVE",A2:A{e})))*D2:D{e})').number_format=CUR
s.cell(g,4).font=Font(bold=True)
rows=[('  of which Amazon.ca',f'=SUMPRODUCT((BOM!$J$2:$J${last}="Amazon.ca")*ISERROR(SEARCH("ALTERNATIVE",BOM!$A$2:$A${last}))*N(+BOM!$I$2:$I${last}))'),
 ('  of which Kent / Home Depot',f'=SUMPRODUCT(ISNUMBER(SEARCH("Kent",BOM!$J$2:$J${last}))*N(+BOM!$I$2:$I${last}))'),
 ('  of which local / industrial suppliers',f'=SUMPRODUCT((ISNUMBER(SEARCH("supplier",BOM!$J$2:$J${last})))*N(+BOM!$I$2:$I${last}))'),
 ('  of which based on VERIFIED prices (amazon.ca / kent.ca)',f'=SUMPRODUCT((ISNUMBER(SEARCH("Verified amazon.ca price",BOM!$N$2:$N${last}))+ISNUMBER(SEARCH("Kent.ca web price",BOM!$N$2:$N${last}))>0)*ISERROR(SEARCH("ALTERNATIVE",BOM!$A$2:$A${last}))*N(+BOM!$I$2:$I${last}))'),
 ('  of which ESTIMATES / placeholders',f'=SUMPRODUCT((ISNUMBER(SEARCH("estimate",BOM!$N$2:$N${last})))*ISERROR(SEARCH("ALTERNATIVE",BOM!$A$2:$A${last}))*N(+BOM!$I$2:$I${last}))'),
 ('ALTERNATIVE rated harness option (if chosen)',f'=SUMPRODUCT(ISNUMBER(SEARCH("ALTERNATIVE",BOM!$A$2:$A${last}))*N(+BOM!$I$2:$I${last}))')]
for k,(lab,f) in enumerate(rows):
    s.cell(g+1+k,1,lab); s.cell(g+1+k,4,f).number_format=CUR
nr=g+len(rows)+2
for t in ['Totals update automatically when you change "Have it" on the BOM sheet (TRUE = already owned -> Qty to buy 0).',
          'Prices are before tax/shipping (NB HST 15 %). Amazon prices change daily; dates are in the BOM sheet.']:
    s.cell(nr,1,t); nr+=1
s.column_dimensions['A'].width=58; s.column_dimensions['B'].width=8; s.column_dimensions['C'].width=18; s.column_dimensions['D'].width=18
s.freeze_panes='A2'
# Amazon cart
c=wb.create_sheet('Amazon cart')
c.append(['#','Item','Qty to buy','Unit price CAD','Line total','Link','Search link'])
for x in c[1]: x.font=hdr; x.fill=fill
c['I1']='Lists only Amazon.ca rows with Qty to buy > 0 (alternatives excluded). Built with INDEX/MATCH on a hidden helper column (BOM!O) so it works in every LibreOffice and Excel version (FILTER needs LO 24.8+/Excel 365). Updates when Have it changes.'
for k in range(1,121):
    r=k+1
    m=f'MATCH({k},BOM!$O$2:$O${last},0)'
    c.cell(r,1,f'=IF(ISNA({m}),"",{k})')
    c.cell(r,2,f'=IF($A{r}="","",INDEX(BOM!$B$2:$B${last},{m}))')
    c.cell(r,3,f'=IF($A{r}="","",INDEX(BOM!$G$2:$G${last},{m}))')
    c.cell(r,4,f'=IF($A{r}="","",INDEX(BOM!$H$2:$H${last},{m}))').number_format=CUR
    c.cell(r,5,f'=IF($A{r}="","",INDEX(BOM!$I$2:$I${last},{m}))').number_format=CUR
    c.cell(r,6,f'=IF($A{r}="","",IF(INDEX(BOM!$K$2:$K${last},{m})="","",HYPERLINK(INDEX(BOM!$K$2:$K${last},{m}),"open on amazon.ca")))')
    c.cell(r,7,f'=IF($A{r}="","",HYPERLINK(INDEX(BOM!$L$2:$L${last},{m}),"search"))')
c.cell(123,2,'Amazon cart total').font=Font(bold=True)
c.cell(123,5,'=SUM(E2:E121)').number_format=CUR
for col,w in zip('ABCDEFG',[5,60,10,13,13,20,10]): c.column_dimensions[col].width=w
c.freeze_panes='A2'
# Notes
nt=wb.create_sheet('Notes')
notes=[
'The Filament Falcon - BOM notes (built 2026-10-08, America/Moncton)',
'Sources: filament-falcon/docs/SPEC.md; ff-fw/docs/electronics/README.md; ff-cad (branch cad, commit b743d58): frame/BOM.md v0.2, frame/cut_list.md, docs/mechanical/README.md, switch_bom.md, displays.md.',
'Filament method: python trimesh volume V and surface area A of every 3MF (panels: single-material body = black, letter inserts = white; printed parts x quantities from docs/mechanical/README.md). Printed volume = shell + infill x (V - shell), shell = min(V, A x 1.2 mm) (3 perimeters x 0.4 mm, also used as top/bottom skin proxy); infill 25 % (monitor standoffs 50 %, harness hub 45 %). PETG 1.27 g/cm3, +15 % for purge/brim/failed prints.',
f'Filament result: black {fil["black"]["cm3"]:.0f} cm3 -> {fil["black"]["kg"]} kg (buy 6 x 1 kg); white {fil["white"]["cm3"]:.0f} cm3 -> {fil["white"]["kg"]} kg (buy 1 x 1 kg). Red/grey parts (guard lids, throttle grip, gear arm, adapter plates) are counted in black - print them in red/grey if you have it. PLA would need ~2 % less mass, but PETG is required for the harness hub and recommended everywhere.',
'Prices: "Verified amazon.ca price" = read from the amazon.ca product page on 2026-10-08 (buy-box price, CAD, before tax). Kent prices are from kent.ca (store shown online: Edmundston); Moncton store prices may differ. Rows marked Estimate/placeholder are NOT verified prices.',
'Motion system (v0.2, final per commit b743d58): 2 cradle actuators 10"/1500 N (~364 mm retracted), 2 pod actuators 12"/1000 N (~415 mm retracted), both self-locking; 2 gas struts ~590 N; 25 mm shafts, UCP205/UCF205/SHF25; spring pawls + latch plates + Bowden release; steel hard stops; key, E-stop, 20 A fuse, hold-to-run DPDT, relays as H-bridges, 10 roller limit switches, safety edge. Check actuator retracted length and self-locking before ordering.',
'SAFETY: the B0B3J1GCVQ buckles are plastic luggage buckles (~50 kg quoted break load, not safety-rated). Positioning harness only for a slow reclining seat. A rated metal-buckle 4-point harness is listed as an ALTERNATIVE row (excluded from totals; switch if you prefer).',
'Displays: GPU DP1 -> C275B, HDMI -> C275W (HDMI/VGA only), DP2 -> DP-to-DVI-D -> TD2220 (VGA/DVI only), USB-A -> USB-B for touch. GPU USB-C is spare (4th display max).',
'Owned items: PC, both Sceptres, TD2220, 3 HY300, 2 joysticks have Have it = TRUE. Buttons/switches/LEDs you may own are FALSE - tick them.',
]
for t in notes: nt.append([t])
nt.column_dimensions['A'].width=160
for row in nt.iter_rows():
    row[0].alignment=Alignment(wrap_text=True,vertical='top')
nt['A1'].font=Font(bold=True,size=13)
wb.save('Filament_Falcon_BOM_src.xlsx')
json.dump(stats,open('stats.json','w')); print(stats, 'rows',n)
